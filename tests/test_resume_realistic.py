"""Realistic resumes from different kinds of job seekers, through the full parse -> skills path.

Each case lists skills the resume plainly shows and skills a careless matcher
would invent from everyday words, hobbies or contact details.
"""

import docx
import pytest

from src.resume.parse_resume import parse_resume
from src.scoring.analyze import resume_skill_names

RESUMES = {
    "psw": (
        """Maria Santos
Brampton, ON | maria.s@example.com

PROFILE
Compassionate Personal Support Worker with 4 years in long-term care.

WORK EXPERIENCE
Personal Support Worker, Maple Grove Long-Term Care
Mar 2022 - Present
- Assist residents with bathing, toileting and feeding; record vital signs.
- Dementia care using Gentle Persuasive Approaches; follow IPAC protocols.
- Meet with families weekly and exchange updates with the nursing team.

CERTIFICATIONS
PSW certificate, CPR Level C, Standard First Aid, WHMIS

HOBBIES
Gardening, hiking, knitting, cooking for my church group
""",
        {"Personal Care", "Vital Signs", "Dementia Care", "Infection Prevention and Control",
         "Personal Support Worker Certificate", "CPR", "First Aid", "WHMIS"},
        {"Cooking", "Landscaping", "Agriculture", "Google Meet", "Microsoft Exchange"},
    ),
    "welder": (
        """Dev Patel
Hamilton, ON

SUMMARY
Red Seal welder, 6 years in structural steel fabrication.

EXPERIENCE
Welder / Fabricator, Steeltown Fabrication
2019 - Present
- MIG and TIG welding of structural steel; blueprint reading and shop drawings.
- Operate forklift and overhead crane; working at heights certified.

SKILLS
MIG, TIG, flux core, hand and power tools, WHMIS, lockout/tagout

INTERESTS
Wood carving, fishing, painting miniatures
""",
        {"Welding", "Red Seal", "Blueprint Reading", "Forklift Operation", "Working at Heights",
         "Power Tools", "WHMIS", "Lockout/Tagout"},
        {"Painting", "Carpentry"},
    ),
    "career_changer": (
        """Sam Rivera
Toronto, ON

SUMMARY
Former line cook moving into data analytics; framing kitchen problems as data questions.

EXPERIENCE
Line Cook, Harbour Bistro
Jun 2018 - Aug 2023
- Prep and cooking on the grill line; Smart Serve certified; food safety.

Data Analytics Student, Seneca Polytechnic
Sep 2023 - Present
- Projects in Python, SQL and Tableau; weekly reports for a capstone client.

SKILLS
Python, SQL, Tableau, Microsoft Excel
""",
        {"Python", "SQL", "Tableau", "Microsoft Excel", "Smart Serve", "Food Safety", "Cooking"},
        {"Carpentry", "Oracle Reports"},
    ),
    "accountant": (
        """Priya Shah, CPA
Mississauga, ON

EXPERIENCE
Staff Accountant, Contoso Bank
Jan 2021 - Present
- Month-end close, journal entries and bank reconciliations.
- Accounts payable and collections; supported the external audit.
- Filing of HST returns; payroll processing in Ceridian.

SKILLS
QuickBooks Online, Sage 50, Microsoft Excel, IFRS

REFERENCES
Available upon request
""",
        {"Month-End Close", "General Ledger", "Account Reconciliation", "Accounts Payable",
         "Accounts Receivable", "Auditing", "Tax Preparation", "Payroll", "QuickBooks",
         "Sage Accounting", "Microsoft Excel", "IFRS", "CPA Designation"},
        set(),
    ),
    "developer": (
        """Jordan Kim
github.com/jkim

EXPERIENCE
Junior Software Developer, Litware
May 2024 - Present
- Built React and TypeScript front ends over a Django REST API.
- Wrote unit tests with pytest; Docker deployment to AWS.
- Low-level assembly language course project; code inspection and debugging.
- Dispatch events with Redux; data labeling tools for an ML team.

SKILLS
Python, JavaScript, TypeScript, React, Django, Docker, AWS, Git
""",
        {"Python", "JavaScript", "TypeScript", "React", "Django", "Docker", "Amazon Web Services", "Git"},
        {"Assembly", "Quality Control", "Dispatching", "Packaging", "Clinical Documentation"},
    ),
}


@pytest.mark.parametrize("name", list(RESUMES))
def test_realistic_resume_skills(name, tmp_path):
    text, expected, invented = RESUMES[name]
    path = tmp_path / f"{name}.txt"
    path.write_text(text, encoding="utf-8")
    skills = set(resume_skill_names(parse_resume(path)))
    assert expected <= skills, f"missing: {expected - skills}"
    assert not invented & skills, f"invented: {invented & skills}"


def test_docx_resume_with_hobbies_table_ignores_hobbies(tmp_path):
    document = docx.Document()
    document.add_paragraph("Taylor Brooks")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, SQL, Power BI")
    document.add_paragraph("HOBBIES & INTERESTS")
    table = document.add_table(rows=1, cols=3)
    for cell, hobby in zip(table.rows[0].cells, ["Wood carving", "Landscaping", "Painting"]):
        cell.text = hobby
    path = tmp_path / "resume.docx"
    document.save(path)

    skills = set(resume_skill_names(parse_resume(path)))
    assert {"Python", "SQL", "Power BI"} <= skills
    assert not {"Landscaping", "Painting", "Carpentry"} & skills


def test_resume_without_any_headings_still_matches(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("I have three years of Python and SQL work, plus Tableau dashboards for sales.", encoding="utf-8")
    assert {"Python", "SQL", "Tableau"} <= set(resume_skill_names(parse_resume(path)))


def test_empty_resume_text_has_no_skills():
    assert resume_skill_names("") == []
