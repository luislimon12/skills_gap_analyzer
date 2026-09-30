"""Role 2 tests: DOCX/PDF parsing, section splitting, experience dates."""

import io
from datetime import date

import docx
import pytest

from src.resume.experience_dates import extract_experience_entries, parse_month, total_months
from src.resume.extract_sections import SECTION_KEYS, extract_sections
from src.resume.parse_resume import ResumeParseError, parse_resume

TODAY = date(2026, 9, 30)

SAMPLE_RESUME = """Jordan Lee
jordan@example.com | 416-555-0199

SUMMARY
Data analyst focused on retail forecasting.

Work Experience
Data Analyst, Acme Retail
Jan 2024 – Present
• Built Tableau dashboards used by 40 stores
• Automated SQL reports

Analytics Intern | Northwind
May 2023 - Aug 2023
- Cleaned data in Python (pandas)

EDUCATION
Seneca Polytechnic — Computer Programming, 2021 - 2023

Technical Skills: Python, SQL, Tableau, Excel

P R O J E C T S
Sales forecaster — scikit-learn
"""


def _make_docx(tmp_path):
    document = docx.Document()
    document.add_paragraph("Jordan Lee")
    document.add_paragraph("SKILLS")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "SQL"
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Data Analyst, Acme Retail — Jan 2024 – Present, building Tableau dashboards for stores")
    path = tmp_path / "resume.docx"
    document.save(path)
    return path


# ---------- parsing ----------

def test_parse_docx_includes_tables_in_document_order(tmp_path):
    text = parse_resume(_make_docx(tmp_path))
    assert text.splitlines()[:3] == ["Jordan Lee", "SKILLS", "Python | SQL"]
    assert text.splitlines()[3] == "EXPERIENCE"


def test_parse_resume_accepts_file_like_upload(tmp_path):
    path = _make_docx(tmp_path)
    upload = io.BytesIO(path.read_bytes())
    upload.name = "My Resume.DOCX"  # Streamlit's UploadedFile exposes .name like this
    assert "Python | SQL" in parse_resume(upload)


def test_parse_pdf(tmp_path):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    path = tmp_path / "resume.pdf"
    c = canvas.Canvas(str(path))
    for i, line in enumerate(["Jordan Lee", "SKILLS", "Python, SQL, Tableau and Excel for reporting work"]):
        c.drawString(72, 750 - 20 * i, line)
    c.save()
    assert parse_resume(path).splitlines() == [
        "Jordan Lee", "SKILLS", "Python, SQL, Tableau and Excel for reporting work"]


def test_unsupported_extension():
    with pytest.raises(ResumeParseError, match="Unsupported file type"):
        parse_resume(io.BytesIO(b"x"), "resume.doc")


def test_corrupt_file_gives_friendly_error():
    with pytest.raises(ResumeParseError, match="corrupt"):
        parse_resume(io.BytesIO(b"not really a pdf"), "resume.pdf")


def test_nearly_empty_document_is_rejected(tmp_path):
    document = docx.Document()
    document.add_paragraph("Jordan")
    path = tmp_path / "empty.docx"
    document.save(path)
    with pytest.raises(ResumeParseError, match="No readable text"):
        parse_resume(path)


# ---------- sections ----------

def test_extract_sections():
    sections = extract_sections(SAMPLE_RESUME)
    assert set(sections) == set(SECTION_KEYS)
    assert sections["contact"].startswith("Jordan Lee")
    assert sections["summary"] == "Data analyst focused on retail forecasting."
    assert sections["experience"].startswith("Data Analyst, Acme Retail")
    assert "Seneca Polytechnic" in sections["education"]
    assert sections["skills"] == "Python, SQL, Tableau, Excel"   # inline "Header: content"
    assert sections["projects"] == "Sales forecaster — scikit-learn"  # spaced-out header
    assert sections["certifications"] == ""


def test_sentence_starting_with_header_word_is_not_a_header():
    sections = extract_sections("EXPERIENCE\nExperience with Python and SQL in production\n")
    assert sections["experience"] == "Experience with Python and SQL in production"


# ---------- experience dates ----------

@pytest.mark.parametrize("text, expected", [
    ("Jan 2022", (2022, 1)), ("September 2021", (2021, 9)), ("Sept. 2023", (2023, 9)),
    ("03/2020", (2020, 3)), ("2019", (2019, 6)), ("13/2020", None),
])
def test_parse_month(text, expected):
    assert parse_month(text) == expected


def test_extract_experience_entries():
    experience = extract_sections(SAMPLE_RESUME)["experience"]
    entries = extract_experience_entries(experience, today=TODAY)
    assert [e["header"] for e in entries] == [
        "Data Analyst, Acme Retail | Jan 2024 – Present",
        "Analytics Intern | Northwind | May 2023 - Aug 2023",
    ]
    first, second = entries
    assert (first["start"], first["end"], first["is_current"], first["months"]) == ("2024-01", "2026-09", True, 33)
    assert "Automated SQL reports" in first["text"]
    assert "pandas" not in first["text"]
    assert second["months"] == 4
    assert "pandas" in second["text"]


def test_entry_with_title_and_dates_on_one_line():
    entries = extract_experience_entries("Barista, Cafe Co  2019 - 2021\n- Customer service", today=TODAY)
    assert len(entries) == 1
    assert entries[0]["months"] == 25  # year-only dates count from June to June


def test_ranges_that_are_not_dates_are_ignored():
    assert extract_experience_entries("Grew revenue from $2000 - 3000 per month", today=TODAY) == []
    assert extract_experience_entries("Shipped 2000 - 3000 units a week", today=TODAY) == []


def test_total_months_merges_overlapping_jobs():
    entries = [
        {"start": "2022-01", "end": "2022-12"},   # 12 months
        {"start": "2022-06", "end": "2023-03"},   # overlaps -> extends to 15 total
        {"start": "2024-01", "end": "2024-02"},   # separate, +2
    ]
    assert total_months(entries) == 17
