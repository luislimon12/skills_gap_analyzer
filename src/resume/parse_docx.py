"""
Role 2 - Component D: extract text from an uploaded DOCX resume.

Splitting the text into sections happens in extract_sections.py. Use
parse_resume.parse_resume() rather than calling this directly.
"""

from typing import BinaryIO

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph


def _table_lines(table: Table) -> list[str]:
    """One line per row. Merged cells repeat in python-docx, so drop consecutive duplicates."""
    lines = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = "\n".join(p.text for p in cell.paragraphs if p.text.strip()).strip()
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return lines


def parse_docx(file: str | BinaryIO) -> str:
    """Return paragraph and table text from a DOCX resume, in document order.

    Many resume templates put the skills list or job headers in tables, which
    document.paragraphs skips, so body elements are walked directly.
    """
    document = docx.Document(file)
    lines = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(element, document).text
            if text.strip():
                lines.append(text)
        elif tag == "tbl":
            lines.extend(_table_lines(Table(element, document)))
    return "\n".join(lines)


if __name__ == "__main__":
    # Smoke test - point this at a sample resume in test_resumes/
    sample = "test_resumes/sample_resume_2.docx"
    print(parse_docx(sample)[:500])
